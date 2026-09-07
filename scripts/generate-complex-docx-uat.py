#!/usr/bin/env python3
"""Create deterministic, fully synthetic complex DOCX UAT fixtures."""

from __future__ import annotations

import csv
import io
import shutil
import sys
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


REPOSITORY = Path(__file__).resolve().parents[1]
KIT = REPOSITORY / "docs" / "acceptance" / "STANDALONE_COMPLEX_DOCX_TEST_KIT"
INPUTS = KIT / "inputs"
FIXED_TIME = (2026, 1, 1, 0, 0, 0)


@dataclass(frozen=True)
class Fixture:
    name: str
    size: str
    scenario: str
    title: str
    person: str = ""
    organisation: str = ""
    email: str = ""
    phone: str = ""
    iban: str = ""
    address: str = ""
    preserve: str = ""
    classification_label: bool = False


SHARED = dict(
    person="Laura Stein",
    organisation="Nordlicht Digital GmbH",
    email="laura.stein@beispiel-firma.de",
    phone="+49 40 555 0182",
    iban="DE89 3704 0044 0532 0130 00",
    address="Wiesenweg 17, 20457 Hamburg",
)

FIXTURES = [
    Fixture("01-kundenprofil-kurz.docx", "kurz", "pii", "Kundenprofil", **SHARED,
            preserve="Product Owner;Markdown;ISO 27001"),
    Fixture("02-projektstatus-kurz.docx", "kurz", "pii", "Projektstatus", **SHARED,
            preserve="Sprint Review;Qualitaetssicherung;Tauri"),
    Fixture("03-technische-notiz-kurz.docx", "kurz", "neutral", "Techniknotiz",
            preserve="Kubernetes;JSON;UTF-8;Retry-Strategie"),
    Fixture("04-betriebskonzept-kurz.docx", "kurz", "neutral", "Betriebskonzept",
            preserve="Offline-Betrieb;Dateisystem;Wiederaufnahme"),
    Fixture("05-label-metadaten-kurz.docx", "kurz", "pii", "Freigabevermerk", **SHARED,
            preserve="VERTRAULICH;Klassifizierungsmetadaten;Dokumentenpruefung",
            classification_label=True),
    Fixture("06-rahmenvertrag-mittel.docx", "mittel", "pii", "Rahmenvertrag",
            person="Murat Kaya", organisation="Elbwiese Beratung GmbH",
            email="murat.kaya@elbwiese-beispiel.de", phone="+49 30 555 0104",
            iban="DE44 5001 0517 5407 3249 31", address="Beispielallee 8, 10115 Berlin",
            preserve="Leistungsbeschreibung;Abnahme;Service Level"),
    Fixture("07-bewerbungsprofil-mittel.docx", "mittel", "pii", "Bewerbungsprofil",
            person="Sofia Lindner", organisation="Morgenrot Systeme KG",
            email="sofia.lindner@morgenrot-beispiel.de", phone="+49 89 555 0147",
            iban="DE12 1002 0500 0001 2345 67", address="Teststrasse 24, 80331 Muenchen",
            preserve="Java;SQL;FHIR;ISTQB;Scrum Master"),
    Fixture("08-prozesshandbuch-mittel.docx", "mittel", "neutral", "Prozesshandbuch",
            preserve="Eingangskontrolle;Vier-Augen-Prinzip;Rollback;Monitoring"),
    Fixture("09-architekturentscheidung-mittel.docx", "mittel", "neutral", "Architekturentscheidung",
            preserve="Separation of Concerns;Sidecar;IPC;Fail Closed"),
    Fixture("10-kundenkorrespondenz-mittel.docx", "mittel", "pii", "Kundenkorrespondenz", **SHARED,
            preserve="Meilenstein;Liefergegenstand;Risikoliste"),
    Fixture("11-fallakte-lang.docx", "lang", "pii", "Fallakte", **SHARED,
            preserve="Sachstand;Massnahmenplan;Entscheidungsprotokoll"),
    Fixture("12-auditbericht-lang.docx", "lang", "pii", "Auditbericht",
            person="Omar Yilmaz", organisation="Hansewerk Digital GmbH",
            email="omar.yilmaz@hansewerk-beispiel.de", phone="+49 69 555 0199",
            iban="DE21 2004 0000 0123 4567 00", address="Pruefweg 5, 60311 Frankfurt",
            preserve="Kontrollziel;Wirksamkeitspruefung;Feststellung"),
    Fixture("13-qualitaetsbericht-lang.docx", "lang", "neutral", "Qualitaetsbericht",
            preserve="Testabdeckung;Regression;Grenzwertanalyse;Fehlertoleranz"),
    Fixture("14-systemhandbuch-lang.docx", "lang", "neutral", "Systemhandbuch",
            preserve="Parser;Konverter;Transaktion;Pruefsumme;Recovery"),
    Fixture("15-projektchronik-lang.docx", "lang", "pii", "Projektchronik", **SHARED,
            preserve="Discovery;Prototyp;Pilotierung;Produktivsetzung"),
]


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def set_cell_text(cell, value: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.size = Pt(9)


def add_header_footer(document: Document, fixture: Fixture) -> None:
    for section in document.sections:
        header = section.header.paragraphs[0]
        header.text = "DATENSECURE UAT | TESTDATEN – VOLLSTAENDIG FIKTIV"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(31, 117, 91)
        footer = section.footer.paragraphs[0]
        footer.text = f"{fixture.title} | Nicht fuer reale Verarbeitung"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(90, 99, 95)


def add_identity_table(document: Document, fixture: Fixture) -> None:
    values = [
        ("Name", fixture.person),
        ("Unternehmen", fixture.organisation),
        ("E-Mail", fixture.email),
        ("Telefon", fixture.phone),
        ("IBAN", fixture.iban),
        ("Anschrift", fixture.address),
    ]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    repeat_table_header(table.rows[0])
    for index, title in enumerate(("Feld", "Fiktiver Testwert")):
        set_cell_text(table.rows[0].cells[index], title, True)
        shade(table.rows[0].cells[index], "DDEFE7")
    for key, value in values:
        cells = table.add_row().cells
        set_cell_text(cells[0], key, True)
        set_cell_text(cells[1], value)


def add_neutral_table(document: Document, fixture: Fixture) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    repeat_table_header(table.rows[0])
    for index, title in enumerate(("Pruefschritt", "Soll", "Status")):
        set_cell_text(table.rows[0].cells[index], title, True)
        shade(table.rows[0].cells[index], "DDEFE7")
    for number, step in enumerate(fixture.preserve.split(";"), start=1):
        cells = table.add_row().cells
        set_cell_text(cells[0], f"T-{number:02d}")
        set_cell_text(cells[1], step)
        set_cell_text(cells[2], "fachlich geprueft")


def chart_bytes(fixture: Fixture) -> bytes:
    image = Image.new("RGB", (1200, 420), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((40, 40, 1160, 380), outline=(31, 117, 91), width=5)
    draw.text((80, 78), "SYNTHETISCHE TESTGRAFIK", fill=(16, 34, 27), font=ImageFont.load_default())
    labels = fixture.preserve.split(";")[:4]
    for index, label in enumerate(labels):
        left = 90 + index * 270
        height = 80 + index * 55
        draw.rectangle((left, 340 - height, left + 150, 340), fill=(31, 117, 91))
        draw.text((left, 350), label[:18], fill=(16, 34, 27), font=ImageFont.load_default())
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=False)
    return buffer.getvalue()


def add_page(document: Document, fixture: Fixture, page: int, total: int) -> None:
    heading = document.add_heading(f"{page}. {fixture.title}: Abschnitt {page}", level=1)
    if page > 1:
        # A break-before heading does not create an empty intermediate page when
        # the previous section happens to end exactly at the page boundary.
        heading.paragraph_format.page_break_before = True
    intro = (
        "Dieses Dokument ist ausschliesslich ein synthetischer UAT-Testfall. "
        "Es kombiniert Fliesstext, Listen, Tabellen und wiederkehrende Fachbegriffe. "
        "Die Originaldatei muss unveraendert bleiben; die Markdown-Ausgabe muss den "
        "fachlichen Inhalt in nachvollziehbarer Reihenfolge wiedergeben."
    )
    document.add_paragraph(intro)
    if fixture.scenario == "pii":
        document.add_paragraph(
            f"{fixture.person} arbeitet fuer {fixture.organisation}. Rueckfragen gehen an "
            f"{fixture.email} oder {fixture.phone}. Die Korrespondenzadresse lautet "
            f"{fixture.address}; das fiktive Abrechnungskonto ist {fixture.iban}."
        )
        document.add_paragraph(
            f"Im gesamten Stapel muessen {fixture.person} und {fixture.organisation} "
            "jeweils konsistent ersetzt werden. Funktionsbezeichnungen, Technologien und "
            "allgemeine Fachsaetze duerfen dadurch nicht veraendert werden."
        )
    else:
        document.add_paragraph(
            "Dieser Kontrollfall enthaelt absichtlich keine Person, keine Kontaktadresse, "
            "keine Bankverbindung und kein kundenspezifisches Unternehmen. Seine Fachbegriffe "
            "sind keine Identifikatoren und sollen bei der Anonymisierung erhalten bleiben."
        )
    preserved = fixture.preserve.split(";")
    document.add_heading("Fachliche Inhalte", level=2)
    for item in preserved:
        document.add_paragraph(
            f"{item} bleibt als technischer Wert und Erhaltungsanker unveraendert.",
            style="List Bullet",
        )
    document.add_heading("Pruefnotizen", level=2)
    note_count = 2 if page == 1 else 4
    for sequence in range(1, note_count + 1):
        document.add_paragraph(
            f"Pruefnotiz {page}.{sequence}: Die Struktur bleibt stabil, der Inhalt wird lokal "
            "verarbeitet und der Testwert wird nicht als reale Aussage interpretiert. "
            f"Abschnitt {page} von {total} prueft Reihenfolge, Absatzgrenzen und Tabellennaehe."
        )
    if page == 1:
        if fixture.scenario == "pii":
            add_identity_table(document, fixture)
        else:
            add_neutral_table(document, fixture)
    elif page % 2 == 0:
        add_neutral_table(document, fixture)
    else:
        image = io.BytesIO(chart_bytes(fixture))
        document.add_picture(image, width=Inches(6.4))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_docx(fixture: Fixture, target: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(28)
    document.core_properties.title = fixture.title
    document.core_properties.subject = "Vollstaendig synthetischer UAT-Kontrollfall"
    document.core_properties.author = ""
    document.core_properties.created = datetime(2026, 1, 1, tzinfo=timezone.utc)
    document.core_properties.modified = datetime(2026, 1, 1, tzinfo=timezone.utc)

    document.add_heading(fixture.title, 0)
    lead = document.add_paragraph("TESTDATEN – VOLLSTAENDIG FIKTIV")
    lead.runs[0].bold = True
    lead.runs[0].font.color.rgb = RGBColor(31, 117, 91)
    document.add_paragraph(
        f"Szenario: {'fiktive Identifikatoren plus Erhaltungsanker' if fixture.scenario == 'pii' else 'neutraler Erhaltungskontrollfall'} | "
        f"Umfang: {fixture.size}"
    )
    page_count = {"kurz": 2, "mittel": 4, "lang": 7}[fixture.size]
    for page in range(1, page_count + 1):
        add_page(document, fixture, page, page_count)
    if fixture.size == "lang":
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_heading("Anhang: strukturierte Kontrollliste", level=1)
        add_neutral_table(document, fixture)
        document.add_paragraph(
            "Der Anhang prueft Abschnittswechsel, wiederholte Kopf- und Fusszeilen sowie "
            "die stabile Reihenfolge der letzten Tabelleninhalte."
        )
    add_header_footer(document, fixture)

    raw = io.BytesIO()
    document.save(raw)
    deterministic_docx(raw.getvalue(), target, fixture.classification_label)


def deterministic_docx(source: bytes, target: Path, classification_label: bool) -> None:
    with zipfile.ZipFile(io.BytesIO(source), "r") as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    if classification_label:
        relationships = members["_rels/.rels"].decode("utf-8")
        marker = "</Relationships>"
        relation = (
            '<Relationship Id="rIdDataSecureLabel" '
            'Type="http://schemas.microsoft.com/office/2020/02/relationships/classificationlabels" '
            'Target="docMetadata/LabelInfo.xml"/>'
        )
        members["_rels/.rels"] = relationships.replace(marker, relation + marker).encode("utf-8")
        content_types = members["[Content_Types].xml"].decode("utf-8")
        override = (
            '<Override PartName="/docMetadata/LabelInfo.xml" '
            'ContentType="application/vnd.ms-office.classificationlabels+xml"/>'
        )
        members["[Content_Types].xml"] = content_types.replace("</Types>", override + "</Types>").encode("utf-8")
        members["docMetadata/LabelInfo.xml"] = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<LabelInfo xmlns="http://schemas.microsoft.com/office/2020/02/metadata/ClassificationLabels">'
            b'<Label id="synthetic-uat-label" enabled="1" setDate="2026-01-01T00:00:00Z"/>'
            b'</LabelInfo>'
        )
    target.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for name in sorted(members):
            info = zipfile.ZipInfo(name, FIXED_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            archive.writestr(info, members[name])


def write_manifest() -> None:
    target = KIT / "EXPECTED_RESULTS.csv"
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow([
            "Datei", "Umfang", "Szenario", "Markdown-Ergebnis", "Anonymisierung",
            "Stapelgruppe", "Muss erhalten bleiben", "Muss ersetzt werden",
        ])
        for fixture in FIXTURES:
            repeated = "Laura-Stein-Nordlicht" if fixture.person == SHARED["person"] else "-"
            redacted = ";".join(filter(None, [fixture.person, fixture.organisation, fixture.email,
                                                   fixture.phone, fixture.iban, fixture.address]))
            writer.writerow([
                fixture.name, fixture.size, fixture.scenario, "1 .md mit gleichem Basisnamen",
                "Identifikatoren ersetzen" if fixture.scenario == "pii" else "Inhalt erhalten",
                repeated, fixture.preserve, redacted or "keine",
            ])


def main() -> int:
    if len(sys.argv) > 2 or (len(sys.argv) == 2 and sys.argv[1] != "--replace"):
        raise SystemExit("Usage: generate-complex-docx-uat.py [--replace]")
    replace = len(sys.argv) == 2
    if INPUTS.exists():
        if not replace:
            raise SystemExit(f"Destination exists: {INPUTS}")
        shutil.rmtree(INPUTS)
    INPUTS.mkdir(parents=True)
    for fixture in FIXTURES:
        build_docx(fixture, INPUTS / fixture.name)
    write_manifest()
    print(f"Created {len(FIXTURES)} deterministic complex DOCX fixtures in {INPUTS}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
